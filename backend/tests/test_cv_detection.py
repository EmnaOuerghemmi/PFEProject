"""
Détection de CV au dépôt.

Quatre familles de garanties :

* le NETTOYAGE ne supprime pas l'information qui fait un CV, masque les
  coordonnées, et reste idempotent — le vectoriseur le rejoue sur des textes
  déjà nettoyés ;
* la VÉRIFICATION DU FICHIER juge le contenu réel, pas le nom ni le type
  déclarés par le navigateur ;
* la DÉCISION applique les seuils (accepté / incertain / refusé / illisible)
  et porte au candidat le message prévu ;
* le MODÈLE SAUVEGARDÉ embarque son nettoyage et ne contient ni email ni
  numéro de téléphone.
"""
import io
import re
import zipfile
from pathlib import Path

import pytest
from fastapi import HTTPException, UploadFile
from httpx import AsyncClient
from starlette.datastructures import Headers

from app.core.config import settings
from app.services import cv_detection, cv_storage
from app.services.cv_detection import (MESSAGE_ILLISIBLE, MESSAGE_REFUS, MOTS_CLES_PRESERVES,
                                       DetecteurCV, VerdictCV, nettoyer_texte)
from app.services.cv_parser import cv_parser

DOSSIER_MODELES = Path(__file__).resolve().parents[1] / "cv_detector" / "models"
MODELE_PRESENT = (DOSSIER_MODELES / DetecteurCV.FICHIER_MODELE).exists()

CV_COURT_EN = """Jane Smith
Software Engineer
EXPERIENCE
Backend Developer, Acme Corp, 2019 - 2023
- Designed and maintained REST APIs in Python and Django
- Managed PostgreSQL databases and Docker deployments
- Mentored two junior developers
EDUCATION
Master in Computer Science, 2019
SKILLS
Python, SQL, Docker, Git, Linux, Django
LANGUAGES
English, French, Arabic"""

CV_COURT_FR = """Sami Ben Salah
Ingénieur logiciel
EXPÉRIENCE PROFESSIONNELLE
Développeur back-end, 2020 - 2024
- Conception d'API REST avec Spring Boot et PostgreSQL
- Mise en place de pipelines CI/CD avec Docker et Jenkins
FORMATION
Diplôme d'ingénieur en informatique, 2020
COMPÉTENCES
Java, Spring Boot, SQL, Docker, Angular, Git
LANGUES
Arabe, français, anglais"""

ATTESTATION_FR = """ATTESTATION DE TRAVAIL
Je soussigné, directeur des ressources humaines de la société, atteste que Monsieur Karim est employé au sein de notre entreprise en qualité de technicien depuis le premier mars 2020.
Cette attestation est délivrée à l'intéressé, à sa demande, pour servir et valoir ce que de droit.
Fait à Tunis, le 12 janvier 2024."""

LETTRE_EN = """Dear Hiring Manager,
I am writing to express my interest in the Data Analyst position at your company. I have always been passionate about data and I believe my background would allow me to contribute to your team.
I would welcome the opportunity to discuss how I can help your organisation. Thank you for your time and consideration.
Sincerely"""


# ── Nettoyage ─────────────────────────────────────────────────────────────────

def test_les_mots_qui_structurent_un_cv_survivent_au_nettoyage():
    texte = nettoyer_texte("Experience Education Skills Formation Compétences Diplôme Attestation")
    assert all(mot in texte.split() for mot in MOTS_CLES_PRESERVES)


def test_les_coordonnees_sont_masquees():
    texte = nettoyer_texte("Contact : jean.test@exemple.com | +216 98 765 432 | linkedin.com/in/jean")
    assert "__email__" in texte and "__telephone__" in texte and "__url__" in texte
    assert "@" not in texte and "765" not in texte


def test_un_numero_decoupe_par_l_extraction_est_masque_aussi():
    """Tabulations et retours à la ligne entre les chiffres : typique d'un PDF."""
    assert nettoyer_texte("Tél : +216\t55 000\n000") == "tel : __telephone__"


def test_une_periode_n_est_pas_un_telephone():
    assert nettoyer_texte("Développeur (2019 - 2022)") == "developpeur ( __annee__ - __annee__ )"


@pytest.mark.parametrize("terme", ["asp.net", "node.js", "15.5/20"])
def test_competences_et_notes_ne_sont_pas_des_adresses_web(terme):
    assert terme in nettoyer_texte(f"Compétences : {terme.upper()} et {terme}")


def test_un_titre_espace_redevient_un_mot():
    assert nettoyer_texte("E X P E R I E N C E  P R O F E S S I O N N E L L E") \
        == "experience professionnelle"


@pytest.mark.parametrize("texte", [
    "a  b  c  d", "Tél : +216\t55 000\n000", "dévelop-\npeur ﬁnance", CV_COURT_FR, ATTESTATION_FR,
])
def test_le_nettoyage_est_idempotent(texte):
    """Le vectoriseur le rejoue sur des textes d'entraînement déjà nettoyés."""
    une_fois = nettoyer_texte(texte)
    assert nettoyer_texte(une_fois) == une_fois


# ── Vérification du fichier ───────────────────────────────────────────────────

PDF_MINIMAL = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n"


def _docx(parties: dict[str, bytes]) -> bytes:
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w") as archive:
        for nom, contenu in parties.items():
            archive.writestr(nom, contenu)
    return tampon.getvalue()


def test_le_type_est_lu_dans_le_contenu():
    assert cv_storage.type_reel(PDF_MINIMAL) == ".pdf"
    assert cv_storage.type_reel(_docx({"word/document.xml": b"<w:document/>"})) == ".docx"


@pytest.mark.parametrize("contenu", [
    b"MZ\x90\x00 un executable renomme en cv.pdf",
    _docx({"word/document.xml": b"<w:document/>", "word/vbaProject.bin": b"macros"}),
    _docx({"xl/workbook.xml": b"<workbook/>"}),
    b"PK\x03\x04 archive corrompue",
])
def test_un_faux_pdf_ou_docx_est_refuse(contenu):
    assert cv_storage.type_reel(contenu) is None


def _depot(contenu: bytes, nom: str, type_mime: str) -> UploadFile:
    return UploadFile(file=io.BytesIO(contenu), filename=nom,
                      headers=Headers({"content-type": type_mime}))


@pytest.mark.asyncio
async def test_le_doc_word_97_n_est_plus_accepte():
    with pytest.raises(HTTPException) as erreur:
        await cv_storage.read_and_validate(_depot(b"\xd0\xcf\x11\xe0", "cv.doc", "application/msword"))
    assert erreur.value.status_code == 400


@pytest.mark.asyncio
async def test_un_fichier_trop_gros_est_refuse_sans_etre_lu_en_entier():
    trop = PDF_MINIMAL + b"0" * cv_storage.MAX_SIZE_BYTES
    with pytest.raises(HTTPException) as erreur:
        await cv_storage.read_and_validate(_depot(trop, "cv.pdf", "application/pdf"))
    assert "5 Mo" in erreur.value.detail


@pytest.mark.asyncio
async def test_l_extension_enregistree_est_celle_du_contenu():
    _, extension = await cv_storage.read_and_validate(_depot(PDF_MINIMAL, "cv.docx", "application/pdf"))
    assert extension == ".pdf"


def test_l_extraction_docx_lit_les_tableaux():
    """Beaucoup de CV Word sont mis en page en tableaux."""
    from docx import Document
    document = Document()
    tableau = document.add_table(rows=1, cols=2)
    tableau.cell(0, 0).text = "COMPÉTENCES"
    tableau.cell(0, 1).text = "Python, SQL"
    tampon = io.BytesIO()
    document.save(tampon)
    texte = cv_parser.extract_text_from_bytes(tampon.getvalue(), ".docx")
    assert "COMPÉTENCES" in texte and "Python, SQL" in texte


# ── Décision ──────────────────────────────────────────────────────────────────

class _DetecteurFixe(DetecteurCV):
    """Détecteur dont la probabilité est imposée : teste les seuils, pas le modèle."""

    def __init__(self, probabilite):
        super().__init__()
        self._probabilite = probabilite

    def probabilite_cv(self, texte):
        return self._probabilite


TEXTE_LISIBLE = "x" * settings.CV_DETECTOR_MIN_TEXT_CHARS


@pytest.mark.parametrize("probabilite, decision, bloquant", [
    (0.95, "accepte", False),
    (0.65, "incertain", False),     # accepté : refuser un vrai CV coûte plus cher
    (0.20, "refuse", True),
])
def test_les_seuils_de_decision(probabilite, decision, bloquant):
    verdict = _DetecteurFixe(probabilite).evaluer(TEXTE_LISIBLE)
    assert verdict.decision == decision and verdict.bloquant is bloquant


def test_un_document_refuse_porte_le_message_destine_au_candidat():
    verdict = _DetecteurFixe(0.10).evaluer(TEXTE_LISIBLE)
    assert verdict.message == MESSAGE_REFUS
    assert "ne semble pas être un CV" in verdict.message


def test_un_document_sans_texte_est_illisible_meme_sans_modele():
    verdict = _DetecteurFixe(None).evaluer("trop court")
    assert verdict.decision == "illisible" and verdict.message == MESSAGE_ILLISIBLE


def test_sans_modele_sur_disque_le_depot_n_est_pas_bloque(tmp_path):
    verdict = DetecteurCV(tmp_path).evaluer(TEXTE_LISIBLE)
    assert verdict.decision == "non_verifie" and not verdict.bloquant


def test_detecteur_desactive(monkeypatch):
    monkeypatch.setattr(settings, "CV_DETECTOR_ENABLED", False)
    assert _DetecteurFixe(0.01).evaluer("").decision == "non_verifie"


@pytest.mark.asyncio
async def test_verifier_document_leve_une_400_avec_le_message(monkeypatch):
    monkeypatch.setattr(cv_detection, "_verifier",
                        lambda contenu, extension: VerdictCV("refuse", 0.1, MESSAGE_REFUS))
    with pytest.raises(HTTPException) as erreur:
        await cv_detection.verifier_document(b"%PDF-", ".pdf")
    assert erreur.value.status_code == 400 and erreur.value.detail == MESSAGE_REFUS


@pytest.mark.asyncio
async def test_verifier_document_laisse_passer_un_cv(monkeypatch):
    monkeypatch.setattr(cv_detection, "_verifier",
                        lambda contenu, extension: VerdictCV("accepte", 0.97))
    assert (await cv_detection.verifier_document(b"%PDF-", ".pdf")).decision == "accepte"


# ── Vérification préalable (formulaire d'inscription) ─────────────────────────
#
# Le compte n'existe pas encore quand le candidat choisit son CV : la route est
# publique, ne stocke rien, et ne renvoie pas la probabilité.

VERIFICATION = "/api/v1/recruitment/cv/verification"


def _envoi(contenu: bytes = PDF_MINIMAL, nom: str = "document.pdf") -> dict:
    return {"cv_file": (nom, io.BytesIO(contenu), "application/pdf")}


@pytest.fixture
def _rien_n_est_stocke(monkeypatch):
    def _interdit(*args, **kwargs):
        raise AssertionError("la vérification préalable ne doit rien écrire sur le disque")
    monkeypatch.setattr(cv_storage, "store_profile_cv", _interdit)
    monkeypatch.setattr(cv_storage, "store_application_cv", _interdit)


@pytest.mark.asyncio
async def test_la_verification_refuse_un_non_cv_sans_compte(client: AsyncClient, monkeypatch,
                                                             _rien_n_est_stocke):
    monkeypatch.setattr(cv_detection, "_verifier",
                        lambda contenu, extension: VerdictCV("refuse", 0.1, MESSAGE_REFUS))
    resp = await client.post(VERIFICATION, files=_envoi())
    assert resp.status_code == 400 and resp.json()["detail"] == MESSAGE_REFUS


@pytest.mark.asyncio
async def test_la_verification_accepte_un_cv_sans_divulguer_la_probabilite(
        client: AsyncClient, monkeypatch, _rien_n_est_stocke):
    monkeypatch.setattr(cv_detection, "_verifier",
                        lambda contenu, extension: VerdictCV("accepte", 0.97))
    resp = await client.post(VERIFICATION, files=_envoi())
    assert resp.status_code == 200 and resp.json() == {"decision": "accepte"}


@pytest.mark.asyncio
async def test_la_verification_controle_aussi_le_type_reel(client: AsyncClient):
    resp = await client.post(VERIFICATION, files=_envoi(b"MZ\x90\x00 executable", "cv.pdf"))
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_la_verification_est_limitee_par_ip(client: AsyncClient, monkeypatch):
    """Route publique qui extrait et classe des documents : sans quota, une
    boucle de requêtes suffirait à occuper le serveur."""
    monkeypatch.setattr(cv_detection, "_verifier",
                        lambda contenu, extension: VerdictCV("accepte", 0.97))
    quota, _ = settings.RATE_LIMIT_CV_CHECK
    for _ in range(quota):
        assert (await client.post(VERIFICATION, files=_envoi())).status_code == 200
    assert (await client.post(VERIFICATION, files=_envoi())).status_code == 429


# ── Modèle sauvegardé ─────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def detecteur_reel():
    if not MODELE_PRESENT:
        pytest.skip("modèle non entraîné (lancer cv_detector/train.py)")
    detecteur = DetecteurCV(DOSSIER_MODELES)
    assert detecteur.disponible
    return detecteur


@pytest.mark.parametrize("texte", [CV_COURT_EN, CV_COURT_FR])
def test_un_cv_court_presente_en_liste_est_accepte(detecteur_reel, texte):
    assert detecteur_reel.evaluer(texte).decision == "accepte"


@pytest.mark.parametrize("texte", [ATTESTATION_FR, LETTRE_EN])
def test_une_attestation_ou_une_lettre_est_refusee(detecteur_reel, texte):
    assert detecteur_reel.evaluer(texte).decision == "refuse"


def test_le_nettoyage_voyage_dans_le_vectoriseur(detecteur_reel):
    """Contrat entraînement / production : impossible de servir un autre nettoyage."""
    assert detecteur_reel._vectoriseur.preprocessor is nettoyer_texte


def test_le_vocabulaire_ne_contient_aucune_coordonnee(detecteur_reel):
    vocabulaire = detecteur_reel._vectoriseur.vocabulary_
    assert not any("@" in terme for terme in vocabulaire)
    assert not any(re.search(r"\d{8,}", terme) for terme in vocabulaire)
    # Termes écartés par min_df : vus dans un seul document — noms, villes…
    assert getattr(detecteur_reel._vectoriseur, "stop_words_", None) is None
