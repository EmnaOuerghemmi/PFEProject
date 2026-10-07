import io
from pathlib import Path
import logging
from typing import BinaryIO, Union

logger = logging.getLogger(__name__)

# Un chemin sur disque, ou un flux en mémoire (dépôt pas encore enregistré).
Source = Union[str, BinaryIO]


class CVParser:
    """
    Extraction du texte d'un CV, PDF ou DOCX.

    C'est la SEULE extraction du projet : l'analyse des CV (`scoring_agent`),
    le détecteur de CV au dépôt (`cv_detection`) et son entraînement
    (`cv_detector/build_dataset.py`) passent tous par elle. Un modèle entraîné
    sur un texte extrait autrement que celui qu'il voit en production apprend
    une distribution qu'il ne retrouvera jamais.
    """

    @staticmethod
    async def extract_text(file_path: str) -> str:
        """Extract text from PDF or DOCX file."""
        try:
            path = Path(file_path)
            if not path.exists():
                logger.warning(f"CV file not found: {file_path}")
                return ""

            if path.suffix.lower() == ".pdf":
                return CVParser._extract_pdf_text(file_path)
            elif path.suffix.lower() in {".docx", ".doc"}:
                return CVParser._extract_docx_text(file_path)
            else:
                logger.warning(f"Unsupported file format: {path.suffix}")
                return ""
        except Exception as e:
            logger.error(f"Error extracting CV text from {file_path}: {str(e)}")
            return ""

    @staticmethod
    def extract_text_from_bytes(contenu: bytes, extension: str) -> str:
        """
        Même extraction, depuis le contenu en mémoire d'un fichier déposé.

        Synchrone et coûteuse (plusieurs centaines de millisecondes pour un
        PDF de quelques pages) : depuis du code asynchrone, l'appeler via
        `asyncio.to_thread`.
        """
        extension = extension.lower()
        if extension == ".pdf":
            return CVParser._extract_pdf_text(io.BytesIO(contenu))
        if extension == ".docx":
            return CVParser._extract_docx_text(io.BytesIO(contenu))
        logger.warning(f"Unsupported file format: {extension}")
        return ""

    @staticmethod
    def _extract_pdf_text(source: Source) -> str:
        """Extract text from PDF using pdfplumber."""
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(source) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            return "\n".join(text_parts)
        except ImportError:
            logger.error("pdfplumber not installed")
            return ""
        except Exception as e:
            logger.error(f"Error extracting PDF text: {str(e)}")
            return ""

    @staticmethod
    def _extract_docx_text(source: Source) -> str:
        """
        Extract text from DOCX using python-docx.

        Paragraphes, mais aussi TABLEAUX et EN-TÊTES : beaucoup de modèles de
        CV Word sont entièrement mis en page en tableaux, et placent nom et
        coordonnées dans l'en-tête. `doc.paragraphs` seul n'en voit rien — un
        tel CV ressortait presque vide, mal noté par l'analyse, et aurait été
        refusé comme illisible par le détecteur.
        """
        try:
            from docx import Document
            text_parts = []
            doc = Document(source)
            for section in doc.sections:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        text_parts.append(para.text)
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    cellules: list[str] = []
                    for cell in row.cells:
                        texte = cell.text.strip()
                        # Une cellule fusionnée est renvoyée une fois par colonne couverte.
                        if texte and texte not in cellules:
                            cellules.append(texte)
                    if cellules:
                        text_parts.append(" | ".join(cellules))
            return "\n".join(text_parts)
        except ImportError:
            logger.error("python-docx not installed")
            return ""
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {str(e)}")
            return ""


cv_parser = CVParser()
